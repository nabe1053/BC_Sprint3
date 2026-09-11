"""xlsx 読取（FUNC-01・AD-004: openpyxl data_only）の単体テスト。

期待インタフェース:
    app.services.readers.xlsx_reader.read_xlsx(data: bytes) -> DocumentReadResult
    - 1シート = 1レコード（locator=シート名, seq=1始まりの出現順）
    - cells には行・列付きのセル値（{"A1": "...", "B2": 123, ...} 形式）
    - data_only 相当で読み、保存済み値の無い数式セルは読めた値として扱わない
      （None のまま or issues に記録。値の捏造をしない）
"""

import io

import openpyxl

from app.core.config import settings
from app.services.readers.xlsx_reader import get_xlsx_sheet_count, read_xlsx
from tests.fixtures.broken_files import (
    make_corrupt_xlsx_bytes,
    make_encrypted_xlsx_bytes,
)

MAX_XLSX_SHEETS = settings.MAX_XLSX_SHEETS


def test_all_sheets_are_extracted_as_separate_records(references_dir) -> None:
    """sample-01（単一シート）: シートが1レコードとして locator=シート名で取れる。"""
    data = (references_dir / "sample-01-tozai-sekiyu-order-list.xlsx").read_bytes()

    result = read_xlsx(data)

    assert result.read_status == "success"
    assert len(result.pages) == 1
    assert result.pages[0].locator == "Order List"
    assert result.pages[0].seq == 1
    assert result.pages[0].cells is not None
    # A1 セルに見出しが入っている
    assert result.pages[0].cells.get("A1") == "オーダーリスト（坑井資材）"


def test_sample02_and_sample08_single_sheet(references_dir) -> None:
    """sample-02 / sample-08 も1シートとして読める（読取成功）。"""
    for name, sheet_name in [
        ("sample-02-hokuyo-energy-shizai-ichiran.xlsx", "所要一覧"),
        ("sample-08-toyosu-kaiyo-hikiai-list.xlsx", "引合リスト"),
    ]:
        data = (references_dir / name).read_bytes()
        result = read_xlsx(data)
        assert result.read_status == "success"
        assert result.pages[0].locator == sheet_name


def test_multi_sheet_workbook_reads_all_sheets_in_order() -> None:
    """複数シートのブックは全シートが対象になり、出現順に seq が振られる。"""
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet-A"
    ws1["A1"] = "alpha"
    ws2 = wb.create_sheet("Sheet-B")
    ws2["A1"] = "beta"
    buf = io.BytesIO()
    wb.save(buf)

    result = read_xlsx(buf.getvalue())

    assert result.read_status == "success"
    assert [p.locator for p in result.pages] == ["Sheet-A", "Sheet-B"]
    assert [p.seq for p in result.pages] == [1, 2]
    assert result.pages[0].cells["A1"] == "alpha"
    assert result.pages[1].cells["A1"] == "beta"


def test_formula_without_cached_value_is_not_treated_as_read() -> None:
    """保存済み値（キャッシュ）が無い数式セルは、読めた値として扱わない
    （捏造しない。値は None のまま／確認事項の材料にする＝issues に記録）。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "label"
    ws["B1"] = '=A1&"!"'  # 数式のみ。保存済み値（キャッシュ）を持たせない
    buf = io.BytesIO()
    wb.save(buf)

    result = read_xlsx(buf.getvalue())

    sheet = result.pages[0]
    assert sheet.cells["A1"] == "label"
    # openpyxl(data_only=True) はキャッシュが無い数式セルに None を返す。
    # 実装はこれを「読めた値」として book に反映してはならない。
    assert sheet.cells.get("B1") is None
    # 中-1: document_issues.locator はシート名と同値にせず、セル座標まで
    # 含めた "{シート名}!{座標}" にする（決定1）。
    issue = next(i for i in result.issues if i.locator == "Sheet1!B1")
    assert issue.issue_type == "unreadable_page"
    # 対象セル座標まで detail に含める（reviewer 指摘 軽微-3: or での曖昧な検証をやめる）
    assert "B1" in issue.detail


def test_corrupt_xlsx_is_unreadable() -> None:
    """zip 構造すら成していない壊れたバイト列は read_status=unreadable。"""
    data = make_corrupt_xlsx_bytes()

    result = read_xlsx(data)

    assert result.read_status == "unreadable"
    assert result.pages == []


def test_xlsx_sheet_count_at_limit_boundary_is_readable() -> None:
    """上限シート数ちょうどのブックも全シート読める（上限境界の読取自体は
    ここでは切り捨てない。上限判定は Service 層の責務。軽微-E: 上限値は
    settings.MAX_XLSX_SHEETS を参照し、ハードコードしない）。"""
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet0"
    for i in range(1, MAX_XLSX_SHEETS):
        wb.create_sheet(f"Sheet{i}")
        wb[f"Sheet{i}"]["A1"] = i
    wb["Sheet0"]["A1"] = 0
    buf = io.BytesIO()
    wb.save(buf)

    result = read_xlsx(buf.getvalue())

    assert result.read_status == "success"
    assert len(result.pages) == MAX_XLSX_SHEETS


def test_all_sheets_empty_is_unreadable_not_success() -> None:
    """全シートが空（値を持つセルが1つも無い）ブックは、読める単位が無いため
    unreadable（契約1・軽微-A）。"""
    wb = openpyxl.Workbook()
    wb.active.title = "EmptySheet"
    wb.create_sheet("AlsoEmpty")
    buf = io.BytesIO()
    wb.save(buf)

    result = read_xlsx(buf.getvalue())

    assert result.read_status == "unreadable"
    assert len(result.issues) == 1
    assert result.issues[0].locator is None


def test_encrypted_xlsx_is_encrypted_not_unreadable() -> None:
    """M-7: 暗号化 .xlsx（OLE2 複合ドキュメント形式）は read_status=encrypted
    として扱う（unreadable に丸めない。PDF と区分をそろえる）。"""
    data = make_encrypted_xlsx_bytes()

    result = read_xlsx(data)

    assert result.read_status == "encrypted"
    assert result.pages == []
    assert len(result.issues) == 1
    assert result.issues[0].issue_type == "encrypted"


def test_page_count_is_sheet_count_when_success() -> None:
    """success 時の page_count はシート数（決定2）。"""
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet-A"
    wb["Sheet-A"]["A1"] = "alpha"
    wb.create_sheet("Sheet-B")
    wb["Sheet-B"]["A1"] = "beta"
    buf = io.BytesIO()
    wb.save(buf)

    result = read_xlsx(buf.getvalue())

    assert result.read_status == "success"
    assert result.page_count == 2


def test_page_count_is_sheet_count_when_all_sheets_empty() -> None:
    """全シート空で unreadable でも、ワークブックとして開けている以上
    シート数は判定できている（決定2: 開けなかった場合のみ None）。"""
    wb = openpyxl.Workbook()
    wb.active.title = "EmptySheet"
    wb.create_sheet("AlsoEmpty")
    buf = io.BytesIO()
    wb.save(buf)

    result = read_xlsx(buf.getvalue())

    assert result.read_status == "unreadable"
    assert result.page_count == 2


def test_page_count_is_none_when_workbook_cannot_be_opened() -> None:
    """破損 xlsx（開けない）は page_count=None のまま（0 で埋めない。中-2）。"""
    data = make_corrupt_xlsx_bytes()

    result = read_xlsx(data)

    assert result.read_status == "unreadable"
    assert result.page_count is None


def test_page_count_is_none_when_encrypted() -> None:
    """暗号化 xlsx も page_count=None のまま（中-2）。"""
    data = make_encrypted_xlsx_bytes()

    result = read_xlsx(data)

    assert result.read_status == "encrypted"
    assert result.page_count is None


def test_get_xlsx_sheet_count_reads_sheet_count_without_extracting_cells() -> None:
    """軽微-2: get_xlsx_sheet_count はセル抽出をせずシート数だけ返す
    （PDF の get_pdf_page_count と対称。上限判定を抽出前に行うため）。"""
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet-A"
    wb.create_sheet("Sheet-B")
    buf = io.BytesIO()
    wb.save(buf)

    assert get_xlsx_sheet_count(buf.getvalue()) == 2


def test_get_xlsx_sheet_count_is_none_for_corrupt_or_encrypted() -> None:
    """軽微-2: 開けない・暗号化の場合は None（0 で埋めない。決定2）。"""
    assert get_xlsx_sheet_count(make_corrupt_xlsx_bytes()) is None
    assert get_xlsx_sheet_count(make_encrypted_xlsx_bytes()) is None


def test_formula_check_opens_workbook_only_once_per_call(monkeypatch) -> None:
    """M-8: 数式判定用のワークブックはシート数に関わらず1回だけ開く
    （N+1 回避。50シート規模でも load_workbook 呼び出しは高々2回:
    値取得用+数式判定用）。"""
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet0"
    wb["Sheet0"]["A1"] = "v0"
    for i in range(1, 10):
        ws = wb.create_sheet(f"Sheet{i}")
        ws["A1"] = f"v{i}"
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()

    load_calls = 0
    original_load_workbook = openpyxl.load_workbook

    def _counting_load_workbook(*args, **kwargs):
        nonlocal load_calls
        load_calls += 1
        return original_load_workbook(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", _counting_load_workbook)

    result = read_xlsx(data)

    assert result.read_status == "success"
    assert load_calls == 2
