"""Read workbook bytes back: projection, no conversion, no formula execution."""
from dataclasses import replace
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace as N
import pytest
from openpyxl import load_workbook
from app.domain.export_types import (
    SHEET_NAMES,
    CASE_HEADERS,
    ITEM_HEADERS,
    EVIDENCE_HEADERS,
    QUESTION_HEADERS,
    RECORD_HEADERS,
    DOCUMENT_HEADERS,
)
from tests.fixtures.export_data import snapshot, AT


def workbook(data=None):
    from app.services.export_workbook import build_workbook

    return load_workbook(
        BytesIO(build_workbook(data or snapshot(), AT)), data_only=False
    )


def cells(wb):
    return [cell for ws in wb for row in ws for cell in row]


def case_values(wb):
    return {row[0].value: row[1].value for row in wb[SHEET_NAMES[0]] if len(row) > 1}


def test_five_sheets_headers_current_values_and_raw_are_preserved():
    wb = workbook()
    assert wb.sheetnames == list(SHEET_NAMES)
    for name, header in zip(
        SHEET_NAMES,
        (
            CASE_HEADERS,
            ITEM_HEADERS,
            EVIDENCE_HEADERS,
            QUESTION_HEADERS,
            RECORD_HEADERS,
        ),
    ):
        actual = tuple(cell.value for cell in next(wb[name].iter_rows())[: len(header)])
        assert actual == header
    ws = wb[SHEET_NAMES[1]]
    assert ws.max_row == 3 and ws.max_column == 55
    assert ws.cell(2, 6).data_type == "n" and Decimal(
        str(ws.cell(2, 6).value)
    ) == Decimal("13.375")
    assert ws.cell(2, 8).value == '13-3/8"'
    assert ws.cell(2, 18).value == "L80" and ws.cell(2, 19).value == "API K55"
    history = wb[SHEET_NAMES[4]]
    assert (
        next(row for row in history.iter_rows(values_only=True) if row[0] == "訂正")[3]
        == "K55"
    )
    assert ws.cell(2, 42).value == 13.375 and ws.cell(2, 46).value == "BOX"
    assert ws.cell(3, 42).value is None


def test_unapproved_conversion_never_creates_other_numeric_values():
    data = snapshot()
    values = {**data.items[0].values}
    for key in values:
        if key.endswith("_value"):
            values[key] = None
    values.update(
        od_value=Decimal("339.7"),
        od_unit="mm",
        od_raw="339.7 mm",
        qty_value=Decimal("118"),
        qty_unit="t",
        qty_raw="118 t",
    )
    data = replace(
        data, items=[replace(data.items[0], values=values, ends=[])], elapsed_sec=None
    )
    numbers = {
        Decimal(str(c.value))
        for c in cells(workbook(data))
        if c.data_type == "n" and c.value is not None
    }
    assert numbers <= {Decimal("339.7"), Decimal("118")}
    assert {Decimal("339.7"), Decimal("118")} <= numbers


@pytest.mark.parametrize(
    "text", ['=HYPERLINK("https://invalid.example","x")', "=1+1", "+1", "-1", "@SUM"]
)
def test_untrusted_strings_are_literal_and_have_no_formula_style_or_link(text):
    data = snapshot()
    data.items[0].values["note"] = text
    data.evidences[0].quote = text
    data.records["unlinked_comments"][0].comment = text
    wb = workbook(data)
    assert wb[SHEET_NAMES[1]].cell(2, 38).value == text
    assert text in [cell.value for cell in cells(wb)]
    for cell in cells(wb):
        assert (
            cell.data_type != "f"
            and cell.hyperlink is None
            and cell.fill.patternType is None
        )
    assert all(len(ws.conditional_formatting) == 0 for ws in wb)


def test_empty_records_have_only_header_and_summary_keeps_unresolved():
    data = snapshot()
    data = replace(data, records={key: [] for key in data.records})
    wb = workbook(data)
    assert wb[SHEET_NAMES[4]].max_row == 1
    values = case_values(wb)
    assert values["未解決件数"] == "1" and values["送付可否"] == "未判断"
    assert values["評価状態"] == "作成案" and values["生成所要（秒）"] == "未記録"


def test_documents_are_received_order_and_snapshot_does_not_cross_cases():
    data = snapshot()
    wb = workbook(data)
    ws = wb[SHEET_NAMES[0]]
    header_row = next(row[0].row for row in ws if row[0].value == "資料名")
    assert (
        tuple(ws.cell(header_row, col).value for col in range(1, 6)) == DOCUMENT_HEADERS
    )
    assert [ws.cell(row, 1).value for row in range(header_row + 1, ws.max_row + 1)] == [
        "synthetic.pdf",
        "later.xlsx",
    ]
    other = snapshot(case=N(id=99, case_code="B_ONLY_CASE"))
    other.items[0].values["note"] = "B_ONLY_ITEM"
    other.documents[0].file_name = "B_ONLY_DOCUMENT"
    foreign = {"B_ONLY_CASE", "B_ONLY_ITEM", "B_ONLY_DOCUMENT"}
    assert foreign & {c.value for c in cells(wb) if isinstance(c.value, str)} == set()


def test_same_snapshot_has_identical_cells_and_fixed_document_properties():
    data = snapshot()
    first = workbook(data)
    second = workbook(data)

    def projection(wb):
        return [
            [
                (cell.coordinate, cell.value, cell.data_type)
                for row in ws
                for cell in row
            ]
            for ws in wb
        ]

    assert projection(first) == projection(second)
    assert first.properties.creator == "OCTG Item List Agent"
    assert (
        first.properties.created == first.properties.modified == AT.replace(tzinfo=None)
    )


def test_history_has_each_record_once_with_kind_ties_and_cancellation_columns():
    from app.domain.export_types import RECORD_KINDS

    data = snapshot()
    data.records["confirmations"].append(
        N(
            id=52,
            kind="row_match",
            item_id=9,
            recorded_by="照合者",
            recorded_at=AT,
            undone_by=None,
            undone_at=None,
        )
    )
    for key, rows in data.records.items():
        for row in rows:
            if key == "bounces":
                row["bounce"].recorded_at = AT
                for comment in row["comments"]:
                    comment.recorded_at = AT
            else:
                row.recorded_at = AT
    ws = workbook(data)[SHEET_NAMES[4]]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert [row[0] for row in rows] == [
        *RECORD_KINDS[:7],
        RECORD_KINDS[6],
        RECORD_KINDS[7],
    ]
    assert rows[0][1:8] == ("R1", "grade", "K55", "記載あり", "L80", "記載あり", "照合")
    assert rows[1][1] == "R1" and rows[1][9] == "照合者"
    assert rows[2][11:] == ("取消者", "2026-09-14 09:00:00+09:00")
    assert rows[3][1:3] == ("R1", "grade") and rows[3][8] == "判断済み/未解決"
    assert rows[4][8] == "評価確認済み→担当者確認済み（未解決 2 件）"
    assert [row[8] for row in rows[6:8]] == ["紐付け済み", "未紐付け"]
    assert rows[8][8] == "承認"
    assert all(row[10] == "2026-09-13 14:03:05+09:00" for row in rows)


def test_elapsed_seconds_and_non_draft_state_are_projected_without_conversion():
    data = snapshot(elapsed_sec=Decimal("3.125"))
    data.version.current_state = "review_checked"
    data.version.is_complete = False
    values = case_values(workbook(data))
    assert Decimal(str(values["生成所要（秒）"])) == Decimal("3.125")
    assert values["評価状態"] == "評価確認済み" and values["一部未完了"] == "あり"
