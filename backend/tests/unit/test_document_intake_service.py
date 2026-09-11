"""DocumentIntakeService（Business Logic層）の単体テスト。Repository は mock する。

期待インタフェース:
    app.services.document_intake_service.DocumentIntakeService(document_repository)
        async def intake_document(
            case_id: int,
            file_name: str,
            file_bytes: bytes,
            received_at: datetime,
            storage_path: str,
        ) -> Document

    流れ: 上限判定（ファイル数／サイズ／PDFページ数／xlsxシート数）
        → 形式判定 → 抽出 → read_status 決定 → repository への保存呼び出し
        （DocumentRepositoryProtocol.create_with_details(document, pages, email_parts,
          issues) -> Document を1回呼ぶ想定）。

    案件内の既存資料件数は Service 自身が
    `document_repository.count_by_case(case_id)` で取得する（上限判定は業務ルール
    のため Business Logic の責務。呼び出し側が数え上げをしない。orchestrator の
    指示により、当初案の `existing_document_count` 引数を廃し、この形に修正した）。

    `storage_path`（原本の保存先パス）は呼び出し側（API 層・T-102）が決定し
    引数で渡す。Service はファイル名を流用せず、受け取った値をそのまま
    `documents.storage_path` に入れる（reviewer 指摘 中-3）。

    app.services.exceptions.LimitExceededError:
        code == "E_LIMIT_EXCEEDED"。details にどの上限を何で超えたかを持つ
        （例: {"limit": "document_count", "max": 50, "actual": 51}）。

    app.services.exceptions.UnsupportedFormatError:
        code == "E_UNSUPPORTED_FORMAT"。05-api-ipo.md 5章 #5:
        投入の事実は記録したうえで送出する（documents に kind='unsupported' /
        read_status='unsupported' で1行残す）。details に
        {"documentId": <保存した資料ID>} を入れる（reviewer 指摘 重大-1）。

D02 上限値（AD-003）は `app.core.config.settings` から取得する（テストの
ハードコードをやめる。reviewer 指摘 軽微-4）。

m-4: 上限値はコンストラクタで注入できる（`DocumentIntakeService(repo,
max_documents_per_case=..., ...)`）。`settings` を書き換えずに境界値テストが
書けることを確認する。

m-5: PDF のページ数上限は全ページ抽出前に判定する（`get_pdf_page_count` で
ページ数だけ取得し、上限超過時は `read_pdf`（全文抽出）を呼ばない）。
"""

import io
from datetime import UTC, datetime
from unittest.mock import AsyncMock

import openpyxl
import pytest
from pypdf import PdfWriter

from app.core.config import settings
from app.services.document_intake_service import DocumentIntakeService
from app.services.exceptions import LimitExceededError, UnsupportedFormatError
from tests.fixtures.broken_files import (
    make_corrupt_pdf_bytes,
    make_corrupt_xlsx_bytes,
    make_encrypted_pdf_bytes,
    make_encrypted_xlsx_bytes,
    make_unsupported_format_bytes,
)

MAX_DOCUMENTS_PER_CASE = settings.MAX_DOCUMENTS_PER_CASE
MAX_FILE_SIZE_MB = settings.MAX_FILE_SIZE_MB
MAX_PDF_PAGES = settings.MAX_PDF_PAGES
MAX_XLSX_SHEETS = settings.MAX_XLSX_SHEETS


def _make_service(
    existing_document_count: int = 0,
    **limit_overrides,
) -> tuple[DocumentIntakeService, AsyncMock]:
    repo = AsyncMock()
    repo.count_by_case.return_value = existing_document_count

    async def _fake_create_with_details(document, *args, **kwargs):
        document.id = 999
        return document

    repo.create_with_details.side_effect = _fake_create_with_details
    return DocumentIntakeService(repo, **limit_overrides), repo


def _pdf_bytes(num_pages: int) -> bytes:
    writer = PdfWriter()
    for _ in range(num_pages):
        writer.add_blank_page(width=100, height=100)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _xlsx_bytes(num_sheets: int) -> bytes:
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet0"
    for i in range(1, num_sheets):
        wb.create_sheet(f"Sheet{i}")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def test_document_count_at_limit_is_processed() -> None:
    """上限ちょうど（既存49件＋今回1件=50件）は切り捨てなく処理される。"""
    service, repo = _make_service(existing_document_count=MAX_DOCUMENTS_PER_CASE - 1)
    data = "本文です。".encode("utf-8")

    await service.intake_document(
        case_id=1,
        file_name="body.txt",
        file_bytes=data,
        received_at=datetime.now(UTC),
        storage_path="/data/cases/1/body.txt",
    )

    repo.create_with_details.assert_awaited_once()


async def test_document_count_exceeding_limit_raises_and_does_not_call_repository() -> (
    None
):
    """51件目の投入は E_LIMIT_EXCEEDED になり、保存 API を一切呼ばない。"""
    service, repo = _make_service(existing_document_count=MAX_DOCUMENTS_PER_CASE)
    data = "本文です。".encode("utf-8")

    with pytest.raises(LimitExceededError) as exc_info:
        await service.intake_document(
            case_id=1,
            file_name="body.txt",
            file_bytes=data,
            received_at=datetime.now(UTC),
            storage_path="/data/cases/1/body.txt",
        )

    assert exc_info.value.code == "E_LIMIT_EXCEEDED"
    assert exc_info.value.details.get("limit") == "document_count"
    repo.create_with_details.assert_not_awaited()


async def test_file_size_at_limit_is_processed() -> None:
    """ファイルサイズちょうど上限（20MB）は切り捨てなく処理される（`>` `>=` の取り違え検出）。"""
    service, repo = _make_service()
    data = b"0" * (MAX_FILE_SIZE_MB * 1024 * 1024)

    await service.intake_document(
        case_id=1,
        file_name="just-at-limit.txt",
        file_bytes=data,
        received_at=datetime.now(UTC),
        storage_path="/data/cases/1/just-at-limit.txt",
    )

    repo.create_with_details.assert_awaited_once()


async def test_file_size_exceeding_limit_raises_and_does_not_call_repository() -> None:
    """ファイルサイズが20MBを超えると E_LIMIT_EXCEEDED。超過内容(limit)が具体的。"""
    service, repo = _make_service()
    oversized = b"0" * (MAX_FILE_SIZE_MB * 1024 * 1024 + 1)

    with pytest.raises(LimitExceededError) as exc_info:
        await service.intake_document(
            case_id=1,
            file_name="huge.txt",
            file_bytes=oversized,
            received_at=datetime.now(UTC),
            storage_path="/data/cases/1/huge.txt",
        )

    assert exc_info.value.code == "E_LIMIT_EXCEEDED"
    assert exc_info.value.details.get("limit") == "file_size"
    repo.create_with_details.assert_not_awaited()


async def test_pdf_page_count_at_limit_is_processed() -> None:
    """PDF 200ページちょうどは切り捨てなく処理される。"""
    service, repo = _make_service()
    data = _pdf_bytes(MAX_PDF_PAGES)

    await service.intake_document(
        case_id=1,
        file_name="200pages.pdf",
        file_bytes=data,
        received_at=datetime.now(UTC),
        storage_path="/data/cases/1/200pages.pdf",
    )

    repo.create_with_details.assert_awaited_once()


async def test_pdf_page_count_exceeding_limit_raises_and_does_not_call_repository() -> (
    None
):
    """PDF 201ページは E_LIMIT_EXCEEDED（超過内容=pdf_pages）。repository を呼ばない。"""
    service, repo = _make_service()
    data = _pdf_bytes(MAX_PDF_PAGES + 1)

    with pytest.raises(LimitExceededError) as exc_info:
        await service.intake_document(
            case_id=1,
            file_name="201pages.pdf",
            file_bytes=data,
            received_at=datetime.now(UTC),
            storage_path="/data/cases/1/201pages.pdf",
        )

    assert exc_info.value.code == "E_LIMIT_EXCEEDED"
    assert exc_info.value.details.get("limit") == "pdf_pages"
    repo.create_with_details.assert_not_awaited()


async def test_xlsx_sheet_count_at_limit_is_processed() -> None:
    """xlsx 50シートちょうどは切り捨てなく処理される（`>` `>=` の取り違え検出）。"""
    service, repo = _make_service()
    data = _xlsx_bytes(MAX_XLSX_SHEETS)

    await service.intake_document(
        case_id=1,
        file_name="50sheets.xlsx",
        file_bytes=data,
        received_at=datetime.now(UTC),
        storage_path="/data/cases/1/50sheets.xlsx",
    )

    repo.create_with_details.assert_awaited_once()


async def test_xlsx_sheet_count_exceeding_limit_raises_and_does_not_call_repository() -> (
    None
):
    """xlsx 51シートは E_LIMIT_EXCEEDED（超過内容=xlsx_sheets）。repository を呼ばない。"""
    service, repo = _make_service()
    data = _xlsx_bytes(MAX_XLSX_SHEETS + 1)

    with pytest.raises(LimitExceededError) as exc_info:
        await service.intake_document(
            case_id=1,
            file_name="51sheets.xlsx",
            file_bytes=data,
            received_at=datetime.now(UTC),
            storage_path="/data/cases/1/51sheets.xlsx",
        )

    assert exc_info.value.code == "E_LIMIT_EXCEEDED"
    assert exc_info.value.details.get("limit") == "xlsx_sheets"
    repo.create_with_details.assert_not_awaited()


async def test_unsupported_format_is_recorded_then_raises_with_document_id() -> None:
    """未対応形式（.pptx）は read_status=unsupported として投入の事実が保存された
    うえで UnsupportedFormatError（E_UNSUPPORTED_FORMAT）を送出する。
    details.documentId に保存済み資料IDが入る（05-api-ipo.md 5章 #5・重大-1）。"""
    service, repo = _make_service()

    with pytest.raises(UnsupportedFormatError) as exc_info:
        await service.intake_document(
            case_id=1,
            file_name="proposal.pptx",
            file_bytes=make_unsupported_format_bytes(),
            received_at=datetime.now(UTC),
            storage_path="/data/cases/1/proposal.pptx",
        )

    assert exc_info.value.code == "E_UNSUPPORTED_FORMAT"
    assert exc_info.value.details.get("documentId") == 999

    repo.create_with_details.assert_awaited_once()
    document = repo.create_with_details.await_args.args[0]
    assert document.kind == "unsupported"
    assert document.read_status == "unsupported"
    pages_arg = repo.create_with_details.await_args.kwargs.get("pages")
    assert not pages_arg


async def test_readable_document_flows_through_to_repository_save() -> None:
    """正常系: 形式判定→抽出→read_status決定→repository保存、という一連の流れを検証する。"""
    service, repo = _make_service()
    data = "見積の本文テキストです。".encode("utf-8")

    result = await service.intake_document(
        case_id=1,
        file_name="body.txt",
        file_bytes=data,
        received_at=datetime.now(UTC),
        storage_path="/data/cases/1/body.txt",
    )

    repo.create_with_details.assert_awaited_once()
    document = repo.create_with_details.await_args.args[0]
    assert document.kind == "text"
    assert document.read_status == "success"
    assert document.storage_path == "/data/cases/1/body.txt"
    assert result is not None


async def test_storage_path_is_used_as_is_not_derived_from_file_name() -> None:
    """storage_path は呼び出し側から渡された値をそのまま使う（file_name を流用しない）。
    reviewer 指摘 中-3。"""
    service, repo = _make_service()
    data = "本文です。".encode("utf-8")

    await service.intake_document(
        case_id=1,
        file_name="report.txt",
        file_bytes=data,
        received_at=datetime.now(UTC),
        storage_path="/storage/2026/09/uuid-1234.txt",
    )

    document = repo.create_with_details.await_args.args[0]
    assert document.storage_path == "/storage/2026/09/uuid-1234.txt"
    assert document.storage_path != document.file_name


async def test_encrypted_pdf_is_not_success_and_records_document_level_issue() -> None:
    """暗号化 PDF は read_status=encrypted で保存され、資料全体（locator=None）の
    document_issues が1件残る（reviewer 指摘 中-2/中-5）。"""
    service, repo = _make_service()

    await service.intake_document(
        case_id=1,
        file_name="secret.pdf",
        file_bytes=make_encrypted_pdf_bytes(),
        received_at=datetime.now(UTC),
        storage_path="/data/cases/1/secret.pdf",
    )

    document = repo.create_with_details.await_args.args[0]
    assert document.read_status == "encrypted"
    assert document.read_status != "success"

    issues_arg = repo.create_with_details.await_args.kwargs.get("issues")
    assert issues_arg
    assert any(i.locator is None and i.issue_type == "encrypted" for i in issues_arg)
    assert all(i.detail for i in issues_arg)


async def test_corrupt_pdf_is_unreadable_and_records_document_level_issue() -> None:
    """破損 PDF は read_status=unreadable で保存され、document_issues が1件残る。"""
    service, repo = _make_service()

    await service.intake_document(
        case_id=1,
        file_name="broken.pdf",
        file_bytes=make_corrupt_pdf_bytes(),
        received_at=datetime.now(UTC),
        storage_path="/data/cases/1/broken.pdf",
    )

    document = repo.create_with_details.await_args.args[0]
    assert document.read_status == "unreadable"
    assert document.read_status != "success"

    issues_arg = repo.create_with_details.await_args.kwargs.get("issues")
    assert issues_arg
    assert any(
        i.locator is None and i.issue_type == "unreadable_page" for i in issues_arg
    )
    assert all(i.detail for i in issues_arg)


async def test_injected_document_count_limit_is_honored_without_touching_settings() -> (
    None
):
    """m-4: コンストラクタで注入した上限値が使われる（settings を書き換えない）。"""
    service, repo = _make_service(existing_document_count=1, max_documents_per_case=1)
    data = "本文です。".encode("utf-8")

    with pytest.raises(LimitExceededError) as exc_info:
        await service.intake_document(
            case_id=1,
            file_name="body.txt",
            file_bytes=data,
            received_at=datetime.now(UTC),
            storage_path="/data/cases/1/body.txt",
        )

    assert exc_info.value.details.get("max") == 1
    repo.create_with_details.assert_not_awaited()
    # settings のグローバル既定値には影響していない。
    assert settings.MAX_DOCUMENTS_PER_CASE == MAX_DOCUMENTS_PER_CASE


async def test_injected_pdf_page_limit_skips_full_extraction_when_exceeded(
    monkeypatch,
) -> None:
    """m-5: 注入した PDF ページ数上限を超える場合、全文抽出（read_pdf）を
    呼ばずにページ数だけで受付拒否する。"""
    from app.services import document_intake_service as module

    read_pdf_calls: list[bytes] = []
    original_read_pdf = module.read_pdf

    def _tracking_read_pdf(data: bytes):
        read_pdf_calls.append(data)
        return original_read_pdf(data)

    monkeypatch.setattr(module, "read_pdf", _tracking_read_pdf)

    service, repo = _make_service(max_pdf_pages=1)
    data = _pdf_bytes(2)

    with pytest.raises(LimitExceededError) as exc_info:
        await service.intake_document(
            case_id=1,
            file_name="2pages.pdf",
            file_bytes=data,
            received_at=datetime.now(UTC),
            storage_path="/data/cases/1/2pages.pdf",
        )

    assert exc_info.value.details.get("limit") == "pdf_pages"
    assert exc_info.value.details.get("max") == 1
    assert read_pdf_calls == []
    repo.create_with_details.assert_not_awaited()


async def test_corrupt_xlsx_is_unreadable_and_records_document_level_issue() -> None:
    """破損 xlsx は read_status=unreadable で保存され、document_issues が1件残る。"""
    service, repo = _make_service()

    await service.intake_document(
        case_id=1,
        file_name="broken.xlsx",
        file_bytes=make_corrupt_xlsx_bytes(),
        received_at=datetime.now(UTC),
        storage_path="/data/cases/1/broken.xlsx",
    )

    document = repo.create_with_details.await_args.args[0]
    assert document.read_status == "unreadable"

    issues_arg = repo.create_with_details.await_args.kwargs.get("issues")
    assert issues_arg
    assert any(
        i.locator is None and i.issue_type == "unreadable_page" for i in issues_arg
    )
    assert all(i.detail for i in issues_arg)


async def test_corrupt_xlsx_page_count_is_none_not_zero() -> None:
    """中-2: 開けなかった xlsx の page_count は None（0 で埋めない）。"""
    service, repo = _make_service()

    await service.intake_document(
        case_id=1,
        file_name="broken.xlsx",
        file_bytes=make_corrupt_xlsx_bytes(),
        received_at=datetime.now(UTC),
        storage_path="/data/cases/1/broken.xlsx",
    )

    document = repo.create_with_details.await_args.args[0]
    assert document.page_count is None


async def test_encrypted_xlsx_page_count_is_none_not_zero() -> None:
    """中-2: 暗号化 xlsx も page_count は None のまま。"""
    service, repo = _make_service()

    await service.intake_document(
        case_id=1,
        file_name="secret.xlsx",
        file_bytes=make_encrypted_xlsx_bytes(),
        received_at=datetime.now(UTC),
        storage_path="/data/cases/1/secret.xlsx",
    )

    document = repo.create_with_details.await_args.args[0]
    assert document.read_status == "encrypted"
    assert document.page_count is None


async def test_text_page_count_is_always_one() -> None:
    """決定2・軽微-5: text の page_count は常に 1。"""
    service, repo = _make_service()
    data = "本文です。".encode("utf-8")

    await service.intake_document(
        case_id=1,
        file_name="body.txt",
        file_bytes=data,
        received_at=datetime.now(UTC),
        storage_path="/data/cases/1/body.txt",
    )

    document = repo.create_with_details.await_args.args[0]
    assert document.page_count == 1


async def test_injected_xlsx_sheet_limit_skips_full_extraction_when_exceeded(
    monkeypatch,
) -> None:
    """軽微-2: 注入した xlsx シート数上限を超える場合、全セル抽出
    （read_xlsx）を呼ばずにシート数だけで受付拒否する。"""
    from app.services import document_intake_service as module

    read_xlsx_calls: list[bytes] = []
    original_read_xlsx = module.read_xlsx

    def _tracking_read_xlsx(data: bytes):
        read_xlsx_calls.append(data)
        return original_read_xlsx(data)

    monkeypatch.setattr(module, "read_xlsx", _tracking_read_xlsx)

    service, repo = _make_service(max_xlsx_sheets=1)
    data = _xlsx_bytes(2)

    with pytest.raises(LimitExceededError) as exc_info:
        await service.intake_document(
            case_id=1,
            file_name="2sheets.xlsx",
            file_bytes=data,
            received_at=datetime.now(UTC),
            storage_path="/data/cases/1/2sheets.xlsx",
        )

    assert exc_info.value.details.get("limit") == "xlsx_sheets"
    assert exc_info.value.details.get("max") == 1
    assert read_xlsx_calls == []
    repo.create_with_details.assert_not_awaited()


async def test_xlsx_success_passes_sheet_cells_through_to_document_pages() -> None:
    """中-4: xlsx の success 経路で pages[].cells が repository に渡される
    ことを検証する（cells → DocumentPage 詰め替えの未検証を解消）。"""
    import openpyxl

    service, repo = _make_service()
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet0"
    wb["Sheet0"]["A1"] = "hello"
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()

    await service.intake_document(
        case_id=1,
        file_name="one-sheet.xlsx",
        file_bytes=data,
        received_at=datetime.now(UTC),
        storage_path="/data/cases/1/one-sheet.xlsx",
    )

    pages_arg = repo.create_with_details.await_args.kwargs.get("pages")
    assert pages_arg
    assert pages_arg[0].locator == "Sheet0"
    assert pages_arg[0].cells == {"A1": "hello"}


async def test_eml_success_passes_part_role_seq_and_sent_at_to_repository() -> None:
    """中-4: .eml の正常系分岐が Service 層でテストされていなかったのを解消。
    email_parts の part_role / seq / sent_at が repository に正しく渡ることを
    検証する。"""
    import email.message

    service, repo = _make_service()
    msg = email.message.EmailMessage()
    msg["From"] = "sender@example.com"
    msg["Subject"] = "quote for OCTG casing"
    msg["Date"] = "Mon, 01 Sep 2026 09:00:00 +0900"
    msg.set_content("見積をお願いします。数量は100本です。")
    data = msg.as_bytes()

    await service.intake_document(
        case_id=1,
        file_name="inquiry.eml",
        file_bytes=data,
        received_at=datetime.now(UTC),
        storage_path="/data/cases/1/inquiry.eml",
    )

    document = repo.create_with_details.await_args.args[0]
    assert document.kind == "eml"
    assert document.read_status == "success"
    assert document.page_count is None

    email_parts_arg = repo.create_with_details.await_args.kwargs.get("email_parts")
    assert email_parts_arg
    latest = next(p for p in email_parts_arg if p.part_role == "latest_body")
    assert latest.seq == 1
    assert latest.sent_at is not None
    assert latest.sent_at.month == 9
    assert latest.sent_at.day == 1
    assert latest.from_addr == "sender@example.com"
    assert "見積をお願いします" in latest.body
