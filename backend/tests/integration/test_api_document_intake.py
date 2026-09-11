"""API #5 資料投入（05-api-ipo.md 1章 A・5章 #5・3章の直接詳細記述は無いが 6章が SSOT）。

対象パス（0.3・区分 UI）: POST /api/v1/ui/cases/{caseId}/documents（multipart/form-data,
フィールド名 "file"）。

このスライスで実装者が新設するもの（orchestrator 決定1・2）:
    app.core.config.settings.STORAGE_ROOT: str  # 既定 "storage"
    保存パス: f"{STORAGE_ROOT}/{case_id}/{uuid4()}{ext}"
    サイズ上限超過はチャンク読取の途中で打ち切り、Service を呼ばずに 413 を返す。

期待するレスポンス契約（仮決め。orval スキーマが SSOT になる前段階）:
    201: {"documentId": <int>, "readStatus": <str>}（reviewer 指摘 軽-6:
         storagePath はサーバ内部パスのため応答に含めない）
    413: {"code": "E_LIMIT_EXCEEDED", "message": ..., "details": {"limit":..., "max":..., "actual":...}}
    415: {"code": "E_UNSUPPORTED_FORMAT", "message": ..., "details": {"documentId": <int>}}
    404: {"code": "E_NOT_FOUND", ...}

reviewer 指摘 中-7: `STORAGE_ROOT` を `tmp_path` に差し替え、実 `backend/storage/`
に書き込まない（`autouse` フィクスチャで全テストに適用する）。
"""

import io

import openpyxl
import pytest
from pypdf import PdfWriter

from app.core.config import settings
from app.repositories.document_repository import DocumentRepository
from tests.fixtures.broken_files import (
    make_encrypted_pdf_bytes,
    make_image_only_pdf_bytes,
    make_unsupported_format_bytes,
)

CASES_PATH = "/api/v1/ui/cases"


@pytest.fixture(autouse=True)
def _use_tmp_storage_root(tmp_path, monkeypatch):
    """中-7: 実 `backend/storage/` に書き込まず、テストごとに使い捨ての
    tmp_path を使う（後始末不要）。"""
    monkeypatch.setattr(settings, "STORAGE_ROOT", str(tmp_path))


def _documents_path(case_id: int) -> str:
    return f"{CASES_PATH}/{case_id}/documents"


def _make_case(client, case_code: str = "CASE-DOC-101") -> int:
    resp = client.post(
        CASES_PATH, json={"caseCode": case_code, "customerName": None, "title": None}
    )
    assert resp.status_code == 201
    return resp.json()["caseId"]


def _make_text_pdf_bytes() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _make_partial_xlsx_bytes() -> bytes:
    """1セルは値、もう1セルは保存済み値の無い数式（読取できないセル）を持つ xlsx。

    xlsx_reader._read_opened_workbook の契約により read_status == 'partial' になる
    （values ありセルが1つ以上、かつ issues が1件以上あるため）。
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws["A2"] = "=A1+1"  # 保存済み値（キャッシュ）を持たない数式
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def test_intake_success_saves_file_under_storage_root_uuid_path(
    client, db_session
) -> None:
    """正常系: 投入したファイルが {STORAGE_ROOT}/{caseId}/{uuid}{ext} に保存され、
    元のファイル名をそのまま流用しない（決定1）。"""
    case_id = _make_case(client)
    file_bytes = _make_text_pdf_bytes()

    response = client.post(
        _documents_path(case_id),
        files={"file": ("original-name.pdf", file_bytes, "application/pdf")},
    )

    assert response.status_code == 201
    body = response.json()
    assert isinstance(body["documentId"], int)
    assert body["readStatus"] in (
        "success",
        "partial",
        "unreadable",
        "encrypted",
        "unsupported",
    )

    repo = DocumentRepository(db_session)
    documents = await repo.list_by_case(case_id)
    saved = next(d for d in documents if d.id == body["documentId"])

    storage_root = settings.STORAGE_ROOT
    assert saved.storage_path.startswith(f"{storage_root}/{case_id}/")
    assert saved.storage_path != "original-name.pdf"
    assert saved.storage_path.endswith(".pdf")

    import os

    assert os.path.isfile(saved.storage_path)
    with open(saved.storage_path, "rb") as f:
        assert f.read() == file_bytes


def test_intake_read_status_is_not_rounded_to_success(client) -> None:
    """正常系: read_status が partial のとき、API は 'success' に丸めずそのまま返す（N03）。"""
    case_id = _make_case(client, case_code="CASE-DOC-PARTIAL")
    file_bytes = _make_partial_xlsx_bytes()

    response = client.post(
        _documents_path(case_id),
        files={
            "file": (
                "partial.xlsx",
                file_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 201
    assert response.json()["readStatus"] == "partial"


def test_intake_unreadable_pdf_returns_unreadable_status(client) -> None:
    """正常系: 画像のみ（テキストが1つも取れない）PDF は unreadable を返す。"""
    case_id = _make_case(client, case_code="CASE-DOC-UNREADABLE")
    file_bytes = make_image_only_pdf_bytes()

    response = client.post(
        _documents_path(case_id),
        files={"file": ("image-only.pdf", file_bytes, "application/pdf")},
    )

    assert response.status_code == 201
    assert response.json()["readStatus"] == "unreadable"


def test_intake_encrypted_pdf_returns_encrypted_status(client) -> None:
    """正常系: 暗号化 PDF は encrypted を返す。"""
    case_id = _make_case(client, case_code="CASE-DOC-ENCRYPTED")
    file_bytes = make_encrypted_pdf_bytes()

    response = client.post(
        _documents_path(case_id),
        files={"file": ("secret.pdf", file_bytes, "application/pdf")},
    )

    assert response.status_code == 201
    assert response.json()["readStatus"] == "encrypted"


def test_intake_exceeding_file_size_limit_returns_413_and_skips_service(
    client, monkeypatch
) -> None:
    """異常系: サイズ上限超過はチャンク読取の途中で打ち切り 413 を返す。
    413 は記録を残さない（documents に行を作らない）（決定2）。"""
    case_id = _make_case(client, case_code="CASE-DOC-413")
    monkeypatch.setattr(settings, "MAX_FILE_SIZE_MB", 0)
    file_bytes = b"x" * 1024  # MAX_FILE_SIZE_MB=0 の下ではどんな内容でも超過

    response = client.post(
        _documents_path(case_id),
        files={"file": ("too-big.pdf", file_bytes, "application/pdf")},
    )

    assert response.status_code == 413
    body = response.json()
    assert body["code"] == "E_LIMIT_EXCEEDED"
    assert "limit" in body["details"]
    assert "max" in body["details"]
    assert "actual" in body["details"]


async def test_intake_exceeding_file_size_limit_does_not_create_document_row(
    client, monkeypatch, db_session
) -> None:
    """異常系: 413 の場合 documents に行が作られない（記録を残さない・決定2）。"""
    case_id = _make_case(client, case_code="CASE-DOC-413B")
    monkeypatch.setattr(settings, "MAX_FILE_SIZE_MB", 0)

    client.post(
        _documents_path(case_id),
        files={"file": ("too-big.pdf", b"x" * 1024, "application/pdf")},
    )

    repo = DocumentRepository(db_session)
    documents = await repo.list_by_case(case_id)
    assert documents == []


def test_intake_unsupported_format_returns_415_with_document_id(client) -> None:
    """異常系: 未対応形式は 415 だが投入の事実は記録され、details.documentId が返る（決定3）。"""
    case_id = _make_case(client, case_code="CASE-DOC-415")
    file_bytes = make_unsupported_format_bytes()

    response = client.post(
        _documents_path(case_id),
        files={"file": ("proposal.pptx", file_bytes, "application/vnd.ms-powerpoint")},
    )

    assert response.status_code == 415
    body = response.json()
    assert body["code"] == "E_UNSUPPORTED_FORMAT"
    assert isinstance(body["details"]["documentId"], int)


async def test_intake_unsupported_format_still_creates_document_row(
    client, db_session
) -> None:
    """異常系: 415 でも documents には kind='unsupported' の行が残る（決定3・AD-005）。"""
    case_id = _make_case(client, case_code="CASE-DOC-415B")
    file_bytes = make_unsupported_format_bytes()

    response = client.post(
        _documents_path(case_id),
        files={"file": ("proposal.pptx", file_bytes, "application/vnd.ms-powerpoint")},
    )
    document_id = response.json()["details"]["documentId"]

    repo = DocumentRepository(db_session)
    documents = await repo.list_by_case(case_id)
    saved = next(d for d in documents if d.id == document_id)
    assert saved.kind == "unsupported"
    assert saved.read_status == "unsupported"


def test_intake_document_for_missing_case_returns_404(client) -> None:
    """異常系: 存在しない案件への投入は 404 E_NOT_FOUND。"""
    response = client.post(
        _documents_path(999999),
        files={"file": ("x.pdf", _make_text_pdf_bytes(), "application/pdf")},
    )

    assert response.status_code == 404
    assert response.json()["code"] == "E_NOT_FOUND"
