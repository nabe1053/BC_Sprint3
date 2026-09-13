"""Real DI, PostgreSQL and isolated disk preservation through export HTTP APIs."""
from datetime import datetime
from hashlib import sha256
from io import BytesIO
from pathlib import Path

import httpx
import pytest
from fastapi import FastAPI
from openpyxl import load_workbook
from sqlalchemy import event
from app.api.errors import ApiError, api_error_handler
from app.api.ui.router import router
from app.core.database import get_db
from app.domain.export_types import SHEET_NAMES
from app.models import Evidence
from tests.fixtures.export_data import seed_export_version


@pytest.fixture
async def export_client(db_session, tmp_path, monkeypatch):
    from app.core.config import settings

    monkeypatch.setattr(settings, "EXPORT_ROOT", str(tmp_path / "exports"))
    app = FastAPI()
    app.add_exception_handler(ApiError, api_error_handler)
    app.include_router(router, prefix="/api/v1")

    async def session():
        yield db_session

    app.dependency_overrides[get_db] = session
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        yield client


async def test_export_http_preserves_exact_workbook_and_initial_integrity(
    db_session, export_client, tmp_path
):
    seed = await seed_export_version(db_session)
    path = f"/api/v1/ui/versions/{seed.version.id}/exports"
    first = await export_client.post(path)
    second = await export_client.post(path)
    assert first.status_code == second.status_code == 200
    assert first.headers["x-export-id"] != second.headers["x-export-id"]
    response = await export_client.get(path)
    assert response.status_code == 200
    rows = response.json()["exports"]
    assert len(rows) == 2 and [r["isInitial"] for r in rows] == [True, False]
    assert [r["exportId"] for r in rows] == [
        int(first.headers["x-export-id"]),
        int(second.headers["x-export-id"]),
    ]
    assert rows[0]["storagePath"] != rows[1]["storagePath"]
    for row, wire in zip(rows, [first, second], strict=True):
        saved = Path(row["storagePath"])
        assert saved.resolve().is_relative_to(tmp_path)
        assert saved.read_bytes() == wire.content
        assert row["contentHash"] == sha256(wire.content).hexdigest()
        assert row["integrity"] == "intact" and row["stateAtExport"] == "draft"
        assert row["sendoffAtExport"] == "undecided" and row["unresolvedAtExport"] == 1
        assert datetime.fromisoformat(row["exportedAt"]).tzinfo is not None
        assert (
            wire.headers["content-disposition"]
            == f'attachment; filename="{row["fileName"]}"'
        )
        wb = load_workbook(BytesIO(wire.content))
        assert wb.sheetnames == list(SHEET_NAMES)
        item_rows = list(wb[SHEET_NAMES[1]].values)
        assert dict(zip(item_rows[0], item_rows[1], strict=True))["グレード"] == "L80"
        case_rows = {r[0].value: r[1].value for r in wb[SHEET_NAMES[0]]}
        assert case_rows["未解決件数"] == "1"
        assert all(
            c.data_type != "f" for ws in wb for row_cells in ws for c in row_cells
        )
    preserved = Path(rows[0]["storagePath"])
    preserved.write_bytes(first.content + b"changed")
    assert [
        r["integrity"] for r in (await export_client.get(path)).json()["exports"]
    ] == ["modified", "intact"]
    preserved.unlink()
    assert [
        r["integrity"] for r in (await export_client.get(path)).json()["exports"]
    ] == ["missing", "intact"]


async def test_bulk_evidence_scope_order_and_query_count_for_two_and_twelve_rows(
    db_session, export_client
):
    seeds = [await seed_export_version(db_session, count=n) for n in (2, 12)]
    for index, seed in enumerate(seeds):
        seed.document.file_name = f"synthetic-case-{index}.pdf"
        for item in reversed(seed.items):
            db_session.add(
                Evidence(
                    version_id=seed.version.id,
                    item_id=item.id,
                    field="grade",
                    raw_value="K55",
                    adopted_value="K55",
                    document_id=seed.document.id,
                    locator=f"row {item.seq}",
                    quote=f"CASE-{index}-ONLY",
                )
            )
    await db_session.commit()
    counts = []
    for index, seed in enumerate(seeds):
        queries = []

        def observe(conn, cursor, statement, parameters, context, executemany):
            if statement.lstrip().upper().startswith("SELECT"):
                queries.append(statement)

        event.listen(db_session.bind.sync_engine, "before_cursor_execute", observe)
        try:
            response = await export_client.get(
                f"/api/v1/ui/versions/{seed.version.id}/evidence"
            )
        finally:
            event.remove(db_session.bind.sync_engine, "before_cursor_execute", observe)
        assert response.status_code == 200
        rows = response.json()["evidences"]
        assert [r["itemId"] for r in rows] == [None, *[item.id for item in seed.items]]
        assert {r["fileName"] for r in rows} == {f"synthetic-case-{index}.pdf"}
        assert f"CASE-{1-index}-ONLY" not in response.text
        assert rows[0]["evidenceId"] == seed.evidence.id
        counts.append(len(queries))
    assert counts == [2, 2]


@pytest.mark.parametrize(
    "method,suffix", [("POST", "exports"), ("GET", "exports"), ("GET", "evidence")]
)
@pytest.mark.parametrize("missing", [False, True])
async def test_pending_and_missing_versions_are_distinct(
    db_session, export_client, method, suffix, missing
):
    seed = await seed_export_version(db_session, finalized=False)
    version_id = 999999999 if missing else seed.version.id
    response = await export_client.request(
        method, f"/api/v1/ui/versions/{version_id}/{suffix}"
    )
    if missing or method == "POST":
        assert response.status_code == (404 if missing else 409)
        assert response.json()["code"] == (
            "E_NOT_FOUND" if missing else "E_VERSION_NOT_FINALIZED"
        )
        assert "x-export-id" not in response.headers
    else:
        assert response.status_code == 200
        if suffix == "exports":
            assert response.json() == {"exports": []}
        else:
            assert response.json()["evidences"][0]["evidenceId"] == seed.evidence.id
