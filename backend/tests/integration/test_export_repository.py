"""Consistent exported snapshots using the same lock as human records."""
import asyncio
from decimal import Decimal
from datetime import timedelta
import pytest
from sqlalchemy import event, text
from app.domain.draft_errors import DraftError
from tests.fixtures.export_data import seed_export_version, AT
from tests.conftest import TestSessionLocal


async def test_snapshot_current_values_and_all_materials_are_version_scoped(db_session):
    from app.repositories.export_repository import ExportRepository

    seed = await seed_export_version(db_session)
    await seed_export_version(db_session)
    repo = ExportRepository(db_session)
    async with repo.lock_version(seed.version.id):
        data = await repo.snapshot(seed.version.id)
    assert data.case.id == seed.case.id and data.version.id == seed.version.id
    assert data.header.id == seed.header.id
    assert [r.values["id"] for r in data.items] == [r.id for r in seed.items]
    assert data.items[0].values["grade"] == "L80"
    assert data.items[0].values["od_value"] == Decimal("13.375")
    assert [e.id for e in data.items[0].ends] == [seed.end.id]
    assert data.items[0].row_match.confirmation_id == seed.confirmation.id
    assert data.items[0].edit_count == 1
    assert [r.id for r in data.documents] == [seed.document.id]
    assert [(r.id, r.file_name) for r in data.evidences] == [
        (seed.evidence.id, seed.document.file_name)
    ]
    assert [r["question"].id for r in data.questions] == [seed.question.id]
    assert data.questions[0]["latest"] is None
    assert set(data.records) == {
        "edits",
        "confirmations",
        "judgements",
        "state_events",
        "bounces",
        "unlinked_comments",
        "sendoff_decisions",
    }
    assert [r.id for r in data.records["edits"]] == [seed.edit.id]
    assert data.unresolved_count == data.matched_count == 1
    assert not data.coverage_confirmed and data.elapsed_sec is None
    assert data.rule_version == seed.rule.rule_version and not data.conversion_enabled


@pytest.mark.parametrize(
    "missing,code", [(True, "E_NOT_FOUND"), (False, "E_VERSION_NOT_FINALIZED")]
)
async def test_lock_errors_distinguish_missing_and_unfinalized(
    db_session, missing, code
):
    from app.repositories.export_repository import ExportRepository

    seed = await seed_export_version(db_session, finalized=False)
    with pytest.raises(DraftError) as error:
        async with ExportRepository(db_session).lock_version(
            seed.version.id + 100 if missing else seed.version.id
        ):
            pytest.fail("unavailable version entered lock")
    assert error.value.code == code


async def test_export_count_and_list_order(db_session):
    from app.repositories.export_repository import ExportRepository
    from tests.fixtures.export_data import export_values

    seed = await seed_export_version(db_session)
    repo = ExportRepository(db_session)
    vid = seed.version.id
    async with repo.lock_version(vid):
        assert await repo.count_exports(vid) == 0
        later = await repo.save_export(
            {**export_values(vid), "exported_at": AT + timedelta(seconds=1)}
        )
        first = await repo.save_export(export_values(vid))
        second = await repo.save_export(export_values(vid))
        assert await repo.count_exports(vid) == 3
    assert [r.id for r in await repo.list_exports(vid)] == [
        first.id,
        second.id,
        later.id,
    ]
    with pytest.raises(DraftError) as error:
        await repo.list_exports(vid + 100)
    assert error.value.code == "E_NOT_FOUND"


async def test_snapshot_query_count_does_not_grow_with_items(db_session):
    from app.repositories.export_repository import ExportRepository

    small = await seed_export_version(db_session, count=2)
    large = await seed_export_version(db_session, count=12)
    counts = []
    for seed in (small, large):
        statements = []

        def observe(conn, cursor, statement, parameters, context, executemany):
            if statement.lstrip().upper().startswith("SELECT"):
                statements.append(statement)

        event.listen(db_session.bind.sync_engine, "before_cursor_execute", observe)
        try:
            repo = ExportRepository(db_session)
            async with repo.lock_version(seed.version.id):
                data = await repo.snapshot(seed.version.id)
        finally:
            event.remove(db_session.bind.sync_engine, "before_cursor_execute", observe)
        counts.append(len(statements))
        assert len(data.items) == len(seed.items)
    assert counts[0] == counts[1] and 1 < counts[0] <= 30


async def test_human_edit_waits_for_export_snapshot_commit(db_session, monkeypatch):
    from app.repositories.export_repository import ExportRepository
    from app.repositories.record_repository import RecordRepository
    from app.services.record_service import RecordService
    from tests.fixtures.record_data import edit_data
    from unittest.mock import Mock

    monkeypatch.setattr(
        "app.services.record_service.datetime",
        Mock(now=lambda tz: AT + timedelta(seconds=1)),
    )
    seed = await seed_export_version(db_session)
    vid, item_id = seed.version.id, seed.item.id
    async with TestSessionLocal() as first, TestSessionLocal() as second:
        second_pid = (
            await second.execute(text("SELECT pg_backend_pid()"))
        ).scalar_one()
        repo = ExportRepository(first)
        task = None
        try:
            async with repo.lock_version(vid):
                before = await repo.snapshot(vid)
                task = asyncio.create_task(
                    RecordService(RecordRepository(second)).edit(
                        vid, edit_data(item_id=item_id, new_value="P110")
                    )
                )
                blocked = False
                for _ in range(100):
                    blocked = (
                        await db_session.execute(
                            text("SELECT cardinality(pg_blocking_pids(:pid)) > 0"),
                            {"pid": second_pid},
                        )
                    ).scalar_one()
                    if blocked:
                        break
                    await asyncio.sleep(0.02)
                assert blocked and not task.done()
                assert before.items[0].values["grade"] == "L80"
            await asyncio.wait_for(task, 5)
            async with repo.lock_version(vid):
                after = await repo.snapshot(vid)
            assert after.items[0].values["grade"] == "P110"
            assert before.items[0].values["grade"] == "L80"
        finally:
            if task is not None and not task.done():
                task.cancel()
                await asyncio.gather(task, return_exceptions=True)


async def test_real_service_exports_two_files_with_matching_hash_and_rows(
    db_session, tmp_path
):
    from pathlib import Path
    from openpyxl import load_workbook
    from app.domain.export_types import SHEET_NAMES, sha256_hex
    from app.repositories.export_repository import ExportRepository
    from app.repositories.document_storage import DocumentStorageGateway
    from app.services.export_service import ExportService

    seed = await seed_export_version(db_session)
    storage = DocumentStorageGateway(str(tmp_path))
    service = ExportService(ExportRepository(db_session), storage)
    first = await service.export(seed.version.id)
    second = await service.export(seed.version.id)
    assert first.record.is_initial and not second.record.is_initial
    assert first.record.id != second.record.id
    for result in (first, second):
        path = storage.resolve_readable_path(result.record.storage_path)
        assert path is not None
        assert Path(path).read_bytes() == result.content
        assert sha256_hex(result.content) == result.record.content_hash
        wb = load_workbook(path)
        assert wb.sheetnames == list(SHEET_NAMES)
        assert wb[SHEET_NAMES[1]].max_row - 1 == len(seed.items)
        assert wb[SHEET_NAMES[1]].cell(2, 18).value == "L80"
    assert [r.integrity for r in await service.list_exports(seed.version.id)] == [
        "intact",
        "intact",
    ]
    pending = await seed_export_version(db_session, finalized=False)
    with pytest.raises(DraftError) as error:
        await service.export(pending.version.id)
    assert error.value.code == "E_VERSION_NOT_FINALIZED"


async def test_snapshot_refreshes_cancelled_records_in_reused_session(db_session):
    from app.repositories.export_repository import ExportRepository
    from app.repositories.record_repository import RecordRepository
    from app.services.record_service import RecordService

    seed = await seed_export_version(db_session)
    vid, eid = seed.version.id, seed.edit.id
    repo = ExportRepository(db_session)
    async with repo.lock_version(vid):
        first = await repo.snapshot(vid)
    assert first.items[0].values["grade"] == "L80"
    async with TestSessionLocal() as other:
        await RecordService(RecordRepository(other)).undo_edit(
            vid, eid, {"recorded_by": "取消者"}
        )
    async with repo.lock_version(vid):
        second = await repo.snapshot(vid)
    assert second.items[0].values["grade"] == "K55" and second.items[0].edit_count == 0
