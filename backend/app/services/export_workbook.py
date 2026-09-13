"""Build a literal, unstyled workbook from an already consistent snapshot."""
from datetime import UTC
from decimal import Decimal
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.writer.excel import ExcelWriter
from app.domain.export_types import (
    SHEET_NAMES,
    CASE_HEADERS,
    DOCUMENT_HEADERS,
    ITEM_HEADERS,
    EVIDENCE_HEADERS,
    QUESTION_HEADERS,
    RECORD_HEADERS,
    RECORD_KINDS,
    ITEM_STATE_LABELS,
    VERSION_STATE_LABELS,
    SENDOFF_LABELS,
    RESOLUTION_LABELS,
    STATUS_LABELS,
    CATEGORY_LABELS,
    format_ts,
)


def write_text(cell, value):
    cell.value = value
    cell.data_type = "s"


def _row(ws, values):
    # IDs and counts are labels; only source Decimals become numeric cells.
    number = ws.max_row + 1 if ws.cell(1, 1).value is not None else 1
    for column, value in enumerate(values, 1):
        cell = ws.cell(number, column)
        if isinstance(value, Decimal):
            cell.value = value
        elif value is not None:
            write_text(cell, str(value))


def _get(value, field):
    return getattr(value, field, None)


def _row_codes(snapshot):
    return {item.values["id"]: item.values["row_code"] for item in snapshot.items}


def build_case_sheet(ws, snapshot, exported_at):
    _row(ws, CASE_HEADERS)
    decisions = snapshot.records["sendoff_decisions"]
    latest = max(decisions, key=lambda r: (r.recorded_at, r.id)) if decisions else None
    entries = [
        ("案件ID", snapshot.case.case_code),
        ("版", snapshot.version.version_no),
        ("作成日時", format_ts(snapshot.version.finalized_at)),
        ("一部未完了", "なし" if snapshot.version.is_complete else "あり"),
        ("評価状態", VERSION_STATE_LABELS[snapshot.version.current_state]),
        ("送付可否", SENDOFF_LABELS[latest.decision if latest else "undecided"]),
        ("送付可否の理由", _get(latest, "reason")),
        ("未解決件数", snapshot.unresolved_count),
        ("一致確認", f"{snapshot.matched_count}/{len(snapshot.items)}"),
        ("網羅性確認", "済" if snapshot.coverage_confirmed else "未"),
        ("訂正件数", sum(item.edit_count for item in snapshot.items)),
        (
            "生成所要（秒）",
            snapshot.elapsed_sec if snapshot.elapsed_sec is not None else "未記録",
        ),
        ("規則版", snapshot.rule_version),
        ("換算", "無効（D03 未承認）"),
    ]
    for entry in entries:
        _row(ws, entry)
    for label, field, state in (
        ("照会番号", "inquiry_no", "inquiry_no_state"),
        ("客先名", "customer_name", "customer_name_state"),
        ("要求納期（原文）", "due_raw", "due_state"),
        ("要求納期の粒度", "due_granularity", None),
        ("納期基準", "due_basis", None),
        ("納地（原文）", "place_raw", "place_state"),
        ("受渡条件", "incoterms", "incoterms_state"),
        ("見積期限（原文）", "quote_deadline_raw", None),
        ("見積期限（日時）", "quote_deadline_at", None),
        ("見積期限 TZ", "quote_deadline_tz_state", None),
    ):
        value = _get(snapshot.header, field)
        if field == "quote_deadline_at":
            value = format_ts(value)
        _row(
            ws,
            (
                label,
                value,
                ITEM_STATE_LABELS.get(_get(snapshot.header, state)) if state else None,
            ),
        )
    _row(ws, ("出力日時", format_ts(exported_at)))
    for note in (
        "サンプル用の簡略書式（D01）",
        ".xlsx は書き出した時点の写し。編集内容はアプリに取り込まれません",
        "未承認の換算は行いません",
    ):
        _row(ws, ("注記", note))
    _row(ws, ("資料一覧",))
    _row(ws, DOCUMENT_HEADERS)
    for doc in sorted(snapshot.documents, key=lambda d: (d.received_at, d.id)):
        _row(
            ws,
            (
                doc.file_name,
                doc.kind,
                doc.read_status,
                format_ts(doc.received_at),
                doc.page_count,
            ),
        )


def build_item_sheet(ws, snapshot, exported_at):
    _row(ws, ITEM_HEADERS)
    fields = (
        "row_code",
        "source_no",
        "kind",
        "kind_raw",
        "usage_note",
        "od_value",
        "od_unit",
        "od_raw",
        "od_state",
        "wall_value",
        "wall_unit",
        "wall_raw",
        "wall_state",
        "weight_value",
        "weight_unit",
        "weight_raw",
        "weight_state",
        "grade",
        "grade_raw",
        "grade_state",
        "connection",
        "connection_raw",
        "connection_state",
        "range_class",
        "length_value",
        "length_unit",
        "length_raw",
        "length_state",
        "qty_value",
        "qty_unit",
        "qty_raw",
        "qty_state",
        "qty_reference_note",
        "due_raw",
        "due_state",
        "place_raw",
        "place_state",
        "note",
        "group_code",
        "candidate_label",
        "is_inherit_candidate",
    )
    for item in sorted(snapshot.items, key=lambda r: (r.values["seq"], r.values["id"])):
        values = [
            ITEM_STATE_LABELS.get(item.values.get(f))
            if f.endswith("_state")
            else item.values.get(f)
            for f in fields
        ]
        values[-1] = "あり" if item.values.get("is_inherit_candidate") else "なし"
        ends = {end.side: end for end in item.ends}
        for side in ("end_a", "end_b"):
            values.extend(
                _get(ends.get(side), f)
                for f in ("od_value", "od_unit", "od_raw", "connection", "thread_end")
            )
        values.extend(
            (
                "済" if item.row_match else "未",
                _get(item.row_match, "recorded_by"),
                format_ts(_get(item.row_match, "recorded_at")),
                item.edit_count,
            )
        )
        _row(ws, values)


def build_evidence_sheet(ws, snapshot, exported_at):
    _row(ws, EVIDENCE_HEADERS)
    codes = _row_codes(snapshot)
    seqs = {item.values["id"]: item.values["seq"] for item in snapshot.items}
    for evidence in sorted(
        snapshot.evidences,
        key=lambda e: (e.item_id is not None, seqs.get(e.item_id, 0), e.id),
    ):
        _row(
            ws,
            (
                codes.get(evidence.item_id, "案件"),
                *(
                    _get(evidence, f)
                    for f in (
                        "field",
                        "raw_value",
                        "adopted_value",
                        "file_name",
                        "locator",
                        "quote",
                        "applied_condition",
                        "conversion_note",
                        "change_reason",
                        "prior_value",
                    )
                ),
            ),
        )


def build_question_sheet(ws, snapshot, exported_at):
    _row(ws, QUESTION_HEADERS)
    codes = _row_codes(snapshot)
    for row in sorted(snapshot.questions, key=lambda r: r["question"].id):
        q, latest = row["question"], row["latest"]
        _row(
            ws,
            (
                q.question_code,
                codes.get(q.item_id, "案件"),
                q.target_field,
                CATEGORY_LABELS.get(q.category),
                q.reason,
                q.candidates,
                STATUS_LABELS[latest.status if latest else "open"],
                RESOLUTION_LABELS[latest.resolution if latest else "unresolved"],
                _get(latest, "note"),
                _get(latest, "recorded_by"),
                format_ts(_get(latest, "recorded_at")),
            ),
        )


def build_record_sheet(ws, snapshot, exported_at):
    _row(ws, RECORD_HEADERS)
    codes = _row_codes(snapshot)
    questions = {r["question"].id: r["question"] for r in snapshot.questions}
    entries = []

    def add(
        kind,
        record,
        item_id=None,
        field=None,
        before=None,
        before_state=None,
        after=None,
        after_state=None,
        reason=None,
        extra=None,
    ):
        entries.append(
            (
                record.recorded_at,
                RECORD_KINDS.index(kind),
                record.id,
                (
                    kind,
                    codes.get(item_id, "案件"),
                    field,
                    before,
                    ITEM_STATE_LABELS.get(before_state),
                    after,
                    ITEM_STATE_LABELS.get(after_state),
                    reason,
                    extra,
                    record.recorded_by,
                    format_ts(record.recorded_at),
                    _get(record, "undone_by"),
                    format_ts(_get(record, "undone_at")),
                ),
            )
        )

    for r in snapshot.records["edits"]:
        add(
            "訂正",
            r,
            r.item_id,
            r.field,
            r.old_value,
            r.old_state,
            r.new_value,
            r.new_state,
            r.reason,
        )
    for r in snapshot.records["confirmations"]:
        add("一致確認" if r.kind == "row_match" else "網羅性確認", r, r.item_id)
    for r in snapshot.records["judgements"]:
        question = questions[r.question_id]
        add(
            "判断",
            r,
            question.item_id,
            question.target_field,
            reason=r.note,
            extra=f"{STATUS_LABELS[r.status]}/{RESOLUTION_LABELS[r.resolution]}",
        )
    for r in snapshot.records["state_events"]:
        add(
            "状態遷移",
            r,
            extra=f"{VERSION_STATE_LABELS[r.from_state]}→{VERSION_STATE_LABELS[r.to_state]}（未解決 {r.unresolved_count} 件）",
        )
    for row in snapshot.records["bounces"]:
        add("差し戻し", row["bounce"], reason=row["bounce"].reason)
        for r in row["comments"]:
            add("差し戻しコメント", r, r.item_id, reason=r.comment, extra="紐付け済み")
    for r in snapshot.records["unlinked_comments"]:
        add("差し戻しコメント", r, r.item_id, reason=r.comment, extra="未紐付け")
    for r in snapshot.records["sendoff_decisions"]:
        add("送付可否", r, reason=r.reason, extra=SENDOFF_LABELS[r.decision])
    for _, _, _, values in sorted(entries, key=lambda row: row[:3]):
        _row(ws, values)


def build_workbook(snapshot, exported_at):
    wb = Workbook()
    wb.remove(wb.active)
    for name, builder in zip(
        SHEET_NAMES,
        (
            build_case_sheet,
            build_item_sheet,
            build_evidence_sheet,
            build_question_sheet,
            build_record_sheet,
        ),
        strict=True,
    ):
        builder(wb.create_sheet(name), snapshot, exported_at)
    wb.properties.creator = "OCTG Item List Agent"
    wb.properties.created = wb.properties.modified = exported_at.astimezone(
        UTC
    ).replace(tzinfo=None)
    stream = BytesIO()
    # save_workbook overwrites modified with wall-clock time; ExcelWriter preserves it.
    with ZipFile(stream, "w", ZIP_DEFLATED, allowZip64=True) as archive:
        ExcelWriter(wb, archive).save()
    return stream.getvalue()
