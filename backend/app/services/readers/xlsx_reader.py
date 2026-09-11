"""xlsx 読取（FUNC-01・AD-004: openpyxl data_only）。

この reader の「読めた単位」はセル（値を持つセル1個 = 1単位）。1シート = 1
`document_pages` レコード（locator=シート名, seq=出現順の1始まり）だが、
契約1（読めた単位が0個 → unreadable）の判定はセル単位で行う。

data_only=True で開くため、保存済み値（キャッシュ）が無い数式セルは None に
なる。これを読めた値として book に反映してはならない（値の捏造をしない。
N03）ため、その旨を document_issues 相当の issues に記録する。

document_issues.locator は document_pages.locator（シート名）と同値にしない。
セル単位の問題は `{シート名}!{セル座標}`（例: "Sheet1!B1"）で記録する
（orchestrator 決定・T-101 中-1）。04-db.md 468行の完了条件（documents_pages
から document_issues の範囲を (document_id, locator) で差し引く）が、
シート名を issue の locator に使うとシート全体を走査済みと相殺してしまう
ため。
"""

from __future__ import annotations

import io

import openpyxl

from app.services.readers.types import DocumentReadResult, PageReadResult, ReadIssue

# OLE2 複合ドキュメント形式の署名（暗号化 .xlsx はこの形式で保存される。M-7）。
# zip 形式（通常の .xlsx）はこの署名を持たない。
_OLE2_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def get_xlsx_sheet_count(data: bytes) -> int | None:
    """xlsx のシート数のみを取得する（セル抽出はしない）。

    上限判定（D02: xlsx シート数上限）に使う（軽微-2: PDF の
    `get_pdf_page_count` と同様、抽出前にシート数だけを先に取得し、上限超過
    なら全セル抽出をせずに受付拒否できるようにする）。暗号化・破損等で
    シート数を判定できない場合は None を返す（例外を投げない。決定2:
    判定できない値は None のまま。0 で埋めない）。
    """
    if data[:8] == _OLE2_SIGNATURE:
        return None
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    except Exception:
        return None
    try:
        return len(workbook.sheetnames)
    finally:
        workbook.close()


def read_xlsx(data: bytes) -> DocumentReadResult:
    """xlsx のバイト列から全シートをセル値付きで抽出する。

    page_count（決定2）: ワークブックとして開けた場合はシート数を返す
    （全シートが空でも、開けている以上シート数は判定できているため）。
    ワークブック自体を開けない（破損）・暗号化で開けない場合は None のまま
    にする（0 で埋めない。中-2）。
    """
    if data[:8] == _OLE2_SIGNATURE:
        # M-7: 暗号化 .xlsx（OLE2 複合ドキュメント）を unreadable に丸めず、
        # PDF と区分をそろえて encrypted として扱う。
        return DocumentReadResult(
            read_status="encrypted",
            issues=[
                ReadIssue(
                    locator=None,
                    issue_type="encrypted",
                    detail="パスワード保護等により復号できない（OLE2 複合ドキュメント形式）",
                )
            ],
        )

    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(data), data_only=True, read_only=True
        )
    except Exception:
        return DocumentReadResult(
            read_status="unreadable",
            issues=[
                ReadIssue(
                    locator=None,
                    issue_type="unreadable_page",
                    detail="xlsx として解釈できない（破損ファイルの可能性）",
                )
            ],
        )

    try:
        # 軽微-4: read_only ワークブックは使い終えたら close() する。
        return _read_opened_workbook(workbook, data)
    finally:
        workbook.close()


def _read_opened_workbook(
    workbook: openpyxl.Workbook, data: bytes
) -> DocumentReadResult:
    # ワークブックとして開けた時点でシート数は判定できている（決定2）。
    page_count = len(workbook.sheetnames)

    # M-8: 数式判定用のブックはシートごとにではなく1回だけ開く（N+1 回避）。
    # 開けない場合は数式判定を諦める（キャッシュ値のみで読む）。
    try:
        formula_workbook = openpyxl.load_workbook(
            io.BytesIO(data), data_only=False, read_only=True
        )
    except Exception:
        formula_workbook = None

    try:
        pages: list[PageReadResult] = []
        issues: list[ReadIssue] = []

        for seq, sheet_name in enumerate(workbook.sheetnames, start=1):
            sheet = workbook[sheet_name]
            cells: dict[str, object] = {}

            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    cells[cell.coordinate] = cell.value

            # data_only=True で読んだときに保存済み値（キャッシュ）が無い数式
            # セルは openpyxl が None を返す。read_only モードではセル型を
            # 取得できないため、元の（数式込み）ワークブックを別途照合する。
            unresolved_formula_cells = _find_formula_cells_without_cached_value(
                formula_workbook, sheet_name, cells
            )

            pages.append(PageReadResult(locator=sheet_name, seq=seq, cells=cells))
            for coordinate in unresolved_formula_cells:
                issues.append(
                    ReadIssue(
                        # 中-1: document_issues.locator はシート名と同値にせず
                        # セル座標まで含める（シート名だけだとページ相殺が起きる）。
                        locator=f"{sheet_name}!{coordinate}",
                        issue_type="unreadable_page",
                        detail=(
                            f"セル {coordinate} は数式の保存済み値（キャッシュ）が無く、"
                            "読取できない"
                        ),
                    )
                )

        total_cells = sum(len(p.cells or {}) for p in pages)
        if total_cells == 0:
            # 契約1: 読めた単位（値を持つセル）が0個 → unreadable（軽微-A）。
            return DocumentReadResult(
                read_status="unreadable",
                page_count=page_count,
                pages=pages,
                issues=[
                    ReadIssue(
                        locator=None,
                        issue_type="unreadable_page",
                        detail="全シートが空でデータを読み取れない",
                    )
                ],
            )

        read_status = "partial" if issues else "success"
        return DocumentReadResult(
            read_status=read_status, page_count=page_count, pages=pages, issues=issues
        )
    finally:
        if formula_workbook is not None:
            formula_workbook.close()


def _find_formula_cells_without_cached_value(
    formula_workbook: openpyxl.Workbook | None,
    sheet_name: str,
    resolved_cells: dict[str, object],
) -> list[str]:
    """数式のみで保存済み値を持たないセルの座標一覧を返す。"""
    if formula_workbook is None or sheet_name not in formula_workbook.sheetnames:
        return []

    formula_sheet = formula_workbook[sheet_name]
    coordinates: list[str] = []
    for row in formula_sheet.iter_rows():
        for cell in row:
            is_formula = isinstance(cell.value, str) and cell.value.startswith("=")
            if is_formula and resolved_cells.get(cell.coordinate) is None:
                coordinates.append(cell.coordinate)
    return coordinates
